#!/usr/bin/env python3
import argparse
import io
import json
import os
import re
import sys
import time
import urllib.error
import urllib.request
from pathlib import Path

from loader import Inbox


CATEGORIES = ["BL_COMPARISON", "SI_REQUEST", "INVOICE_QUERY", "GENERAL", "SPAM"]
REQUIRED_FIELDS = [
    "shipper",
    "consignee",
    "notify_party",
    "port_of_loading",
    "port_of_discharge",
    "container_count",
    "gross_weight_kg",
]


VALID_STATUSES = {"OK", "MISMATCH", "NEEDS_REVIEW"}
VALID_REVIEW_REASONS = {None, "wrong_doc_type", "missing_attachment", "unreadable", "missing_value"}


SYSTEM_PROMPT = """You are an expert shipping documentation analyst.

Classify each email into exactly one category:
- BL_COMPARISON: compare Shipping Instruction (SI) against draft Bill of Lading (BL), or a BL draft/checking workflow.
- SI_REQUEST: asks for or provides shipping instruction details, but is not asking to compare SI vs BL.
- INVOICE_QUERY: invoice, billing, charges, GR, cancellation, freight/local-charge query.
- GENERAL: operational, HR, reminder, status, reports, RPA notifications, ordinary non-spam email.
- SPAM: phishing, scam, marketing, suspicious prize/parcel/account email.

For BL_COMPARISON with SI and BL document text, compare exactly these seven fields by meaning:
shipper, consignee, notify_party, port_of_loading, port_of_discharge, container_count, gross_weight_kg.

Return:
- status OK if all seven fields match.
- status MISMATCH if one or more compared fields differ. Include exactly the differing fields.
- status NEEDS_REVIEW only when the comparison cannot be decided because of wrong_doc_type, missing_attachment, unreadable, or missing_value.

For non-BL_COMPARISON categories, use status OK, review_reason null, has_defect false, defect_fields [].
Respond with only one valid JSON object. No markdown, no explanation."""


def provider_name():
    return os.environ.get("AI_PROVIDER", "openai").lower().strip()


def api_endpoint():
    base = os.environ.get("AI_API_BASE", "https://api.openai.com/v1")
    base = base.rstrip("/")
    if base.endswith("/chat/completions"):
        return base
    return base + "/chat/completions"


def api_key():
    if provider_name() == "gemini":
        key = os.environ.get("GEMINI_API_KEY") or os.environ.get("AI_API_KEY")
    else:
        key = os.environ.get("OPENAI_API_KEY") or os.environ.get("AI_API_KEY")
    if not key:
        raise RuntimeError(
            "Missing API key.\n"
            "  OpenAI PowerShell: $env:OPENAI_API_KEY = 'your_key_here'\n"
            "  OpenAI-compatible PowerShell: $env:AI_API_KEY = 'your_key_here'\n"
            "  Gemini PowerShell: $env:GEMINI_API_KEY = 'your_key_here'"
        )
    return key


def model_name():
    if provider_name() == "gemini":
        return os.environ.get("GEMINI_MODEL", os.environ.get("AI_MODEL", "gemini-2.0-flash"))
    return os.environ.get("AI_MODEL", "gpt-4o-mini")


def trim_text(text, limit=12000):
    text = text or ""
    if len(text) <= limit:
        return text
    head = text[: limit // 2]
    tail = text[-limit // 2 :]
    return head + "\n\n[...middle truncated...]\n\n" + tail


def read_attachment(inbox, att_path):
    suffix = Path(att_path).suffix.lower()
    try:
        raw = inbox.read_bytes(att_path)
    except Exception:
        return "", False

    if suffix == ".txt":
        return raw.decode("utf-8", errors="replace"), True

    if suffix == ".xlsx":
        try:
            from openpyxl import load_workbook

            wb = load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
            lines = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
                    if len(cells) >= 2:
                        lines.append(f"{cells[0]}: {cells[1]}")
                    elif cells:
                        lines.append(cells[0])
            return "\n".join(lines), True
        except Exception:
            return "", False

    if suffix == ".docx":
        try:
            from docx import Document

            doc = Document(io.BytesIO(raw))
            lines = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    cells = [c.text.strip() for c in row.cells]
                    if len(cells) >= 2 and cells[0]:
                        lines.append(f"{cells[0]}: {cells[1]}")
            return "\n".join(lines), True
        except Exception:
            return "", False

    if suffix == ".pdf":
        try:
            import pdfplumber

            texts = []
            with pdfplumber.open(io.BytesIO(raw)) as pdf:
                for page in pdf.pages:
                    page_text = page.extract_text() or ""
                    if page_text.strip():
                        texts.append(page_text)
            text = "\n".join(texts)
            return text, bool(text.strip())
        except Exception:
            try:
                try:
                    import pymupdf as fitz
                except Exception:
                    import fitz

                doc = fitz.open(stream=raw, filetype="pdf")
                text = "\n".join(page.get_text() or "" for page in doc)
                return text, bool(text.strip())
            except Exception:
                return "", False

    return "", False


def build_user_prompt(email, attachment_payloads):
    payload = {
        "email_id": email.get("email_id"),
        "from": email.get("from"),
        "subject": email.get("subject"),
        "body": trim_text(email.get("body", ""), 8000),
        "attachments": attachment_payloads,
        "required_output_schema": {
            "category": CATEGORIES,
            "status": ["OK", "MISMATCH", "NEEDS_REVIEW"],
            "review_reason": [None, "wrong_doc_type", "missing_attachment", "unreadable", "missing_value"],
            "has_defect": "boolean",
            "defect_fields": REQUIRED_FIELDS,
        },
    }
    return json.dumps(payload, ensure_ascii=False, indent=2)


def call_ai(prompt, max_retries=3):
    if provider_name() == "gemini":
        return call_gemini(prompt, max_retries)
    return call_openai_compatible(prompt, max_retries)


def call_openai_compatible(prompt, max_retries=3):
    body = {
        "model": model_name(),
        "temperature": 0,
        "messages": [
            {"role": "system", "content": SYSTEM_PROMPT},
            {"role": "user", "content": prompt},
        ],
        "response_format": {"type": "json_object"},
    }
    data = json.dumps(body).encode("utf-8")
    headers = {
        "Authorization": f"Bearer {api_key()}",
        "Content-Type": "application/json",
    }

    last_error = None
    for attempt in range(max_retries):
        req = urllib.request.Request(api_endpoint(), data=data, headers=headers, method="POST")
        try:
            with urllib.request.urlopen(req, timeout=90) as resp:
                result = json.loads(resp.read().decode("utf-8"))
            content = result["choices"][0]["message"]["content"]
            return parse_json_object(content)
        except (urllib.error.HTTPError, urllib.error.URLError, TimeoutError, KeyError, json.JSONDecodeError) as exc:
            last_error = exc
            if isinstance(exc, urllib.error.HTTPError):
                try:
                    detail = exc.read().decode("utf-8", errors="replace")
                    last_error = f"HTTP {exc.code}: {detail}"
                except Exception:
                    last_error = f"HTTP {exc.code}: {exc.reason}"
            if attempt + 1 < max_retries:
                time.sleep(2 * (attempt + 1))
    raise RuntimeError(f"AI request failed after {max_retries} attempts: {last_error}")


def check_openai_connection():
    old_provider = os.environ.get("AI_PROVIDER")
    old_base = os.environ.get("AI_API_BASE")
    os.environ["AI_PROVIDER"] = "openai"
    os.environ.pop("AI_API_BASE", None)
    try:
        result = call_openai_compatible(
            'Return exactly this JSON object: {"ok": true}',
            max_retries=1,
        )
    finally:
        if old_provider is None:
            os.environ.pop("AI_PROVIDER", None)
        else:
            os.environ["AI_PROVIDER"] = old_provider
        if old_base is not None:
            os.environ["AI_API_BASE"] = old_base

    if result.get("ok") is True:
        print(f"OpenAI connection OK. Endpoint: {api_endpoint()} Model: {model_name()}")
        return True
    raise RuntimeError(f"OpenAI connection returned unexpected JSON: {result}")


def show_config():
    key_var = "GEMINI_API_KEY" if provider_name() == "gemini" else "OPENAI_API_KEY"
    fallback_var = "AI_API_KEY"
    has_primary = bool(os.environ.get(key_var))
    has_fallback = bool(os.environ.get(fallback_var))
    endpoint = gemini_endpoint().split("?key=", 1)[0] if provider_name() == "gemini" else api_endpoint()
    print(f"provider        : {provider_name()}")
    print(f"model           : {model_name()}")
    print(f"endpoint        : {endpoint}")
    print(f"{key_var} set : {has_primary}")
    print(f"{fallback_var} set     : {has_fallback}")


def gemini_endpoint():
    base = os.environ.get("GEMINI_API_BASE", "https://generativelanguage.googleapis.com/v1beta")
    base = base.rstrip("/")
    model = model_name()
    if model.startswith("models/"):
        model_path = model
    else:
        model_path = f"models/{model}"
    return f"{base}/{model_path}:generateContent?key={api_key()}"


def call_gemini(prompt, max_retries=3):
    body = {
        "systemInstruction": {
            "parts": [{"text": SYSTEM_PROMPT}],
        },
        "contents": [
            {
                "role": "user",
                "parts": [{"text": prompt}],
            }
        ],
        "generationConfig": {
            "temperature": 0,
            "response_mime_type": "application/json",
        },
    }
    data = json.dumps(body).encode("utf-8")
    headers = {"Content-Type": "application/json"}

    last_error = None
    for attempt in range(max_retries):
        req = urllib.request.Request(gemini_endpoint(), data=data, headers=headers, method="POST")
        try:
            with urllib.request.urlopen(req, timeout=90) as resp:
                result = json.loads(resp.read().decode("utf-8"))
            content = extract_gemini_text(result)
            return parse_json_object(content)
        except (urllib.error.HTTPError, urllib.error.URLError, TimeoutError, KeyError, json.JSONDecodeError) as exc:
            last_error = exc
            if isinstance(exc, urllib.error.HTTPError):
                try:
                    detail = exc.read().decode("utf-8", errors="replace")
                    last_error = f"HTTP {exc.code}: {detail}"
                except Exception:
                    last_error = f"HTTP {exc.code}: {exc.reason}"
            if attempt + 1 < max_retries:
                time.sleep(2 * (attempt + 1))
    raise RuntimeError(f"Gemini request failed after {max_retries} attempts: {last_error}")


def extract_gemini_text(result):
    candidates = result.get("candidates") or []
    if not candidates:
        raise KeyError(f"No Gemini candidates returned: {result}")
    parts = candidates[0].get("content", {}).get("parts", [])
    texts = [part.get("text", "") for part in parts if part.get("text")]
    if not texts:
        raise KeyError(f"No Gemini text returned: {result}")
    return "\n".join(texts)


def parse_json_object(text):
    text = text.strip()
    try:
        return json.loads(text)
    except json.JSONDecodeError:
        match = re.search(r"\{.*\}", text, flags=re.S)
        if not match:
            raise
        return json.loads(match.group(0))


def normalize_ai_row(row):
    category = row.get("category")
    status = row.get("status")
    review_reason = row.get("review_reason")
    defect_fields = row.get("defect_fields") or []
    has_defect = bool(row.get("has_defect"))

    if category not in CATEGORIES:
        category = "GENERAL"
    if status not in VALID_STATUSES:
        status = "OK"
    if review_reason not in VALID_REVIEW_REASONS:
        review_reason = None
    defect_fields = [field for field in defect_fields if field in REQUIRED_FIELDS]

    if category != "BL_COMPARISON":
        return {
            "category": category,
            "status": "OK",
            "review_reason": None,
            "has_defect": False,
            "defect_fields": [],
        }

    if status == "MISMATCH":
        has_defect = True
        review_reason = None
    elif status == "NEEDS_REVIEW":
        has_defect = False
        defect_fields = []
        review_reason = review_reason or "unreadable"
    else:
        has_defect = False
        defect_fields = []
        review_reason = None

    return {
        "category": category,
        "status": status,
        "review_reason": review_reason,
        "has_defect": has_defect,
        "defect_fields": defect_fields,
    }


def attachment_payloads(inbox, email):
    out = []
    for att in email.get("attachments", []):
        text, readable = read_attachment(inbox, att)
        out.append(
            {
                "path": att,
                "readable_text": readable,
                "text": trim_text(text, 14000),
            }
        )
    return out


def run_ai_pipeline(source, output, limit=None):
    inbox = Inbox(source)
    emails = inbox.emails()
    if limit:
        emails = emails[:limit]

    submission = {}
    for index, email in enumerate(emails, start=1):
        eid = email["email_id"]
        prompt = build_user_prompt(email, attachment_payloads(inbox, email))
        row = normalize_ai_row(call_ai(prompt))
        submission[eid] = row
        print(f"{index}/{len(emails)} {eid}: {row['category']} {row['status']}")

    Path(output).write_text(json.dumps(submission, indent=2), encoding="utf-8")
    print(f"Wrote {len(submission)} predictions to {output}")
    return submission


def main():
    parser = argparse.ArgumentParser(description="AI-based SDOC pipeline using OpenAI-compatible APIs or Gemini.")
    parser.add_argument("--source", default="resources/sdoc-hackathon-bundle")
    parser.add_argument("--output", default="submission_ai.json")
    parser.add_argument("--limit", type=int, help="process only the first N emails for testing cost/latency")
    parser.add_argument("--provider", choices=["gemini", "openai"], help="override AI_PROVIDER for this run")
    parser.add_argument("--check-api", action="store_true", help="test the OpenAI API key/model connection and exit")
    parser.add_argument("--show-config", action="store_true", help="print provider/model/endpoint without printing secrets")
    args = parser.parse_args()
    if args.provider:
        os.environ["AI_PROVIDER"] = args.provider
    if args.show_config:
        show_config()
        return
    if args.check_api:
        check_openai_connection()
        return
    run_ai_pipeline(args.source, args.output, args.limit)


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        sys.exit(1)
