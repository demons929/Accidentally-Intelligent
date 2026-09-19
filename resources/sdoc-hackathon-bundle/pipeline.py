import json
import re
from loader import Inbox
<<<<<<< HEAD
=======
import io
from pathlib import Path
>>>>>>> d9b3d766d193e3a6950a6a68cf24e88c15870f18

# ─── Field Aliases & Patterns ──────────────────────────────────────
FIELD_ALIASES = {
    "shipper": ["shipper", "exporter", "consignor", "sender"],
    "consignee": ["consignee", "receiver", "to the order of", "non-negotiable"],
    "notify_party": ["notify party", "notify", "notification"],
    "port_of_loading": ["port of loading", "load port", "loading port", "pol"],
    "port_of_discharge": ["port of discharge", "discharge port", "pod", "destination"],
    "container_count": ["no. of containers", "container count", "total containers", "packages"],
    "gross_weight_kg": ["gross weight", "gross wt", "gw", "total weight"],
}

REQUIRED_FIELDS = list(FIELD_ALIASES.keys())

<<<<<<< HEAD
=======
def extract_text_from_attachment(inbox, att_path):
    """Reads .txt, .xlsx, .docx, and .pdf files and returns plain text."""
    try:
        # 1. Handle Text Files (Original Logic)
        if att_path.endswith(".txt"):
            return inbox.read_text(att_path)
        
        # Get raw bytes for binary files
        raw_bytes = inbox.read_bytes(att_path)
        
        # 2. Handle Excel Files (.xlsx)
        if att_path.endswith(".xlsx"):
            from openpyxl import load_workbook
            wb = load_workbook(filename=io.BytesIO(raw_bytes))
            text_parts = []
            for sheet in wb.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    # Filter out None values and join cells
                    line = " ".join([str(cell) for cell in row if cell is not None])
                    if line.strip():
                        text_parts.append(line)
            return "\n".join(text_parts)

        # 3. Handle Word Files (.docx)
        elif att_path.endswith(".docx"):
            from docx import Document
            doc = Document(io.BytesIO(raw_bytes))
            return "\n".join([para.text for para in doc.paragraphs])

        # 4. Handle PDF Files (.pdf)
        elif att_path.endswith(".pdf"):
            import pdfplumber
            with pdfplumber.open(io.BytesIO(raw_bytes)) as pdf:
                text_parts = []
                for page in pdf.pages:
                    text_parts.append(page.extract_text())
                return "\n".join(text_parts)
                
        else:
            return "" # Unsupported format
            
    except Exception as e:
        print(f"⚠️ Error reading {att_path}: {e}")
        return ""
    
>>>>>>> d9b3d766d193e3a6950a6a68cf24e88c15870f18
def normalize_value(field, value):
    """Clean up extracted values for better comparison."""
    if not value: return ""
    value = str(value).strip().upper()
    
    # Remove common units or noise
    if field == "gross_weight_kg":
        value = re.sub(r'[^\d.]', '', value)
        try: return float(value)
        except: return value
    if field == "container_count":
        nums = re.findall(r'\d+', value)
        return int(nums[0]) if nums else value
    
    # Normalize port names (remove codes in parentheses for comparison)
    if "port" in field:
        value = re.sub(r'\s*\(.*?\)', '', value).strip()
        
    return value

def parse_document(text):
    """Heuristic parser for SI/BL text files."""
    result = {}
    lines = text.split('\n')
    
    for line in lines:
        if ':' not in line: continue
        key, _, val = line.partition(':')
        key_clean = key.strip().lower()
        val_clean = val.strip()
        
        if not val_clean: continue
        
        for canonical, aliases in FIELD_ALIASES.items():
            if any(alias in key_clean for alias in aliases):
                result[canonical] = normalize_value(canonical, val_clean)
                break
                
    return result

def classify_email(email):
    """Classify based on subject/body keywords."""
    text = f"{email.get('subject', '')} {email.get('body', '')}".lower()
    
    if "spam" in text or "verify your account" in text or "weird trick" in text:
        return "SPAM"
    if "invoice" in text and ("query" in text or "breakdown" in text or "charge" in text):
        return "INVOICE_QUERY"
    if "shipping instruction" in text and ("attached" not in text and "find attached" not in text) and "draft bl" not in text:
        # If it's just asking for an SI or providing one without BL comparison
        return "SI_REQUEST"
    if "bl" in text and ("compare" in text or "check" in text or "confirm" in text or "match" in text):
        return "BL_COMPARISON"
    
    # Default logic for emails with both SI and BL attachments
    atts = email.get("attachments", [])
    has_si = any("si" in a.lower() for a in atts)
    has_bl = any("bl" in a.lower() for a in atts)
    if has_si and has_bl:
        return "BL_COMPARISON"
        
    return "GENERAL"

def compare_si_bl(si_text, bl_text):
    si_data = parse_document(si_text)
    bl_data = parse_document(bl_text)
    
    if not si_data or not bl_data:
        return {"status": "NEEDS_REVIEW", "review_reason": "unreadable", "has_defect": False, "defect_fields": []}

    defect_fields = []
    for field in REQUIRED_FIELDS:
        si_val = si_data.get(field)
        bl_val = bl_data.get(field)
        
        if si_val is None or bl_val is None:
            # Check if it's a major missing field
            if field in ["port_of_loading", "port_of_discharge", "consignee"]:
                defect_fields.append(field)
            continue
            
        if si_val != bl_val:
            defect_fields.append(field)
            
    if defect_fields:
        return {"status": "MISMATCH", "review_reason": None, "has_defect": True, "defect_fields": defect_fields}
    
    return {"status": "OK", "review_reason": None, "has_defect": False, "defect_fields": []}
<<<<<<< HEAD
=======

>>>>>>> d9b3d766d193e3a6950a6a68cf24e88c15870f18
def run_pipeline(source="resources/sdoc-hackathon-bundle"):
    print(f"🔍 Initializing Inbox from: {source}")
    try:
        inbox = Inbox(source)
        emails = inbox.emails()
        print(f"✅ Found {len(emails)} emails.")
    except Exception as e:
        print(f"❌ Error loading inbox: {e}")
        return

    submission = {}
    
    for i, email in enumerate(emails):
        eid = email["email_id"]
        print(f"⚙️ Processing {i+1}/{len(emails)}: {eid}")
        
        category = classify_email(email)
        
        entry = {
            "category": category,
            "status": None,
            "review_reason": None,
            "has_defect": False,
            "defect_fields": [],
        }
        
        if category == "BL_COMPARISON":
            si_text = ""
            bl_text = ""
            
            for att in email.get("attachments", []):
<<<<<<< HEAD
                try:
                    if "si" in att.lower():
                        si_text = inbox.read_text(att)
                    if "bl" in att.lower():
                        bl_text = inbox.read_text(att)
                except Exception as e:
                    print(f"   ⚠️ Could not read attachment {att}: {e}")

            # Fallback to body if SI attachment is missing/unreadable
=======
                content = extract_text_from_attachment(inbox, att)
                
                if "si" in att.lower():
                    si_text = content
                if "bl" in att.lower():
                    bl_text = content

            # Fallback: Check email body for SI if no attachment found or failed
>>>>>>> d9b3d766d193e3a6950a6a68cf24e88c15870f18
            if not si_text and "shipper" in email.get("body", "").lower():
                si_text = email["body"]

            if not si_text or not bl_text:
<<<<<<< HEAD
                entry.update({"status": "NEEDS_REVIEW", "review_reason": "missing_attachment"})
=======
                entry.update({"status": "NEEDS_REVIEW", "review_reason": "unreadable"})
>>>>>>> d9b3d766d193e3a6950a6a68cf24e88c15870f18
            else:
                cmp = compare_si_bl(si_text, bl_text)
                entry.update(cmp)
                
        submission[eid] = entry
        
    if submission:
        with open("submission.json", "w") as f:
            json.dump(submission, f, indent=2)
        print(f"🚀 Successfully wrote {len(submission)} entries to submission.json")
    else:
        print("⚠️ Submission dictionary is still empty!")

if __name__ == "__main__":
    run_pipeline()